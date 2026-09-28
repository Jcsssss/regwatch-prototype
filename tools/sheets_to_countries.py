#!/usr/bin/env python3
"""Porter le classeur comparatif case par case, sans l'aplatir.

    python3 tools/sheets_to_countries.py <classeur.xlsx>   # -> src/reg/nis2/data_sheets.js

Pourquoi un second extracteur a cote de excel_to_countries.py
  Celui-la choisit : il projette 136 champs sur une quinzaine de champs types
  plus des puces de prose, ce qui sert la carte, les KPI et les tris. La fiche
  pays refondue a besoin de l'inverse - la structure du classeur telle quelle,
  une section par feuille, une ligne par colonne - pour qu'une case ajoutee au
  Cyber Watch apparaisse sans qu'on retouche une table de correspondance.

  Les deux coexistent donc, et ne se marchent pas dessus : les champs types
  restent la source des tris et des graphiques, celui-ci la source de la fiche.

Les libelles sont ecrits une fois, pas par pays, et les cases vides sont
omises : sans ces deux regles le fichier triplait pour ne porter que du vide.
"""

import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, column_index_from_string
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
# Une carte par reglementation : data/cellmap/<reg>.json (voir tools/cellmap.py).
sys.path.insert(0, str(Path(__file__).resolve().parent))
from cellmap import path_for  # noqa: E402

CELLMAP = path_for("nis2")
OUT = ROOT / "src" / "reg" / "nis2" / "data_sheets.js"

# Nom court par feuille : il sert de cle dans les donnees et d'ancre dans la page.
KEYS = {
    "ID - P1": "id",
    "Incident reporting - P1": "inc",
    "Registration - P1": "reg",
    "Cybersecurity frameworks - P1": "fw",
    "Audit & Controls - P1": "aud",
    "Sanctions - P2": "san",
    "Authority - P3": "auth",
}

# Colonnes superseded par une autre : le classeur porte des paires
# ANCIENNE / NOUVELLE COLONNE, et lire les deux afficherait deux reponses
# contradictoires a la meme question. La nouvelle fait foi.
OLD = re.compile(r"ANCIENNE COLONNE\s*$", re.I)
NEW = re.compile(r"\s*NOUVELLE COLONNE\s*$", re.I)

# Colonnes de service, ecartees quel que soit leur contenu.
#
#   Repartition   qui, dans l'equipe, tient ce pays a jour. Elle porte des
#                 prenoms : c'est de la donnee personnelle, et la page est
#                 servie publiquement par GitHub Pages. Elle n'a rien a faire
#                 dans un fichier distribue, meme si le depot est prive.
#   Column1...    artefact des tableaux Excel, toujours vide.
#   en-tete       un en-tete purement numerique n'est pas un libelle mais une
#   numerique     donnee : la ligne d'en-tete y porte deja une valeur.
SKIP_LABEL = re.compile(r"^(r[ée]partition|column\d*|colonne\d*)$", re.I)
SKIP_NUMERIC = re.compile(r"^[\d\s.,%-]+$")


def clean(v):
    """Une valeur de cellule, ramenee a quelque chose d'affichable."""
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return "YES" if v else "NO"
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).strip()
    # Le classeur charrie des artefacts d'export et des espaces insecables.
    s = s.replace("_x000D_", " ").replace("_x0002_", "").replace("\xa0", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s).strip()
    if not s or s.lower() in ("n/a", "na", "-", "--", "?"):
        return None
    # « Non communique » n'est pas « vide » : le classeur dit qu'on a cherche et
    # qu'on n'a pas trouve. Les trois graphies sont ramenees a une seule, que la
    # fiche affiche comme telle plutot que de la faire passer pour une reponse.
    if re.fullmatch(r"n\.?\s?c\.?|tbc|tbd", s, re.I):
        return "NC"
    return s


def main():
    if len(sys.argv) < 2:
        raise SystemExit("usage: sheets_to_countries.py <classeur.xlsx>")
    src = Path(sys.argv[1]).expanduser()
    if not src.exists():
        raise SystemExit("classeur introuvable : %s" % src)

    cm = json.loads(CELLMAP.read_text(encoding="utf-8"))
    wb = load_workbook(src, data_only=True, read_only=True)

    sheets, rows, kept, dropped = {}, {}, 0, 0

    for name, spec in cm["sheets"].items():
        key = KEYS.get(name)
        if not key:
            continue
        ws = wb[name]
        # Les colonnes sont lues dans la feuille, pas dans la cartographie.
        # Celle-ci vieillit des qu'on insere une colonne dans le classeur - ce
        # qui est arrive : l'ajout de « Information required in the early
        # warning » a decale deux colonnes, et l'extracteur lisait alors le
        # libelle d'une colonne sur les donnees d'une autre. La cartographie ne
        # sert plus qu'a situer la ligne d'en-tete et celles des pays, qui ne
        # bougent pas quand on ajoute une colonne.
        header = spec["headerRow"]
        first = column_index_from_string(spec["countryColumn"]) + 1
        fields = []
        for idx in range(first, (ws.max_column or first) + 1):
            col = get_column_letter(idx)
            label = clean(ws.cell(row=header, column=idx).value)
            # Un en-tete qui est en fait une donnee : certaines colonnes portent
            # une phrase entiere en guise de libelle.
            if not label or len(label) > 90:
                dropped += 1
                continue
            if OLD.search(label) or SKIP_LABEL.match(label) or SKIP_NUMERIC.match(label):
                dropped += 1
                continue
            fields.append({"c": col, "l": NEW.sub("", label).strip()})
        sheets[key] = {"name": name, "fields": fields}

        for iso, row in spec["rows"].items():
            bag = {}
            for f in fields:
                val = clean(ws["%s%d" % (f["c"], row)].value)
                if val is not None:
                    bag[f["c"]] = val
                    kept += 1
            rows.setdefault(iso, {})[key] = bag

    wb.close()
    # La date du passage : c'est elle que la fiche affiche comme date de mise a
    # jour. Le classeur n'a pas de date par pays, et les dates de ses feuilles
    # disent quand quelqu'un l'a edite, pas quand l'outil l'a relu.
    payload = {"workbook": src.name, "scanned": date.today().isoformat(),
               "sheets": sheets, "rows": rows}
    OUT.write_text(
        "/* ---- Le classeur comparatif, porte case par case.\n"
        "   Genere par tools/sheets_to_countries.py - ne pas editer a la main.\n"
        "   Le classeur sur SharePoint est la source de verite. ---- */\n"
        "const SHEET_DATA = %s;\n" % json.dumps(payload, ensure_ascii=False,
                                                separators=(",", ":")),
        encoding="utf-8")

    total = sum(len(s["fields"]) for s in sheets.values())
    print("%d feuilles, %d champs retenus (%d ecartes)" % (len(sheets), total, dropped))
    print("%d cases renseignees sur %d pays" % (kept, len(rows)))
    print("%d Ko -> %s" % (OUT.stat().st_size // 1024, OUT.relative_to(ROOT)))


if __name__ == "__main__":
    main()
