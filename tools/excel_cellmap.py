#!/usr/bin/env python3
"""Build the address book of the comparative KPI workbook.

Every fact RegWatch shows lives in one cell of `CYBER WATCH5_Technical inventory.xlsx`,
addressed by (sheet, country row, field column). This script extracts that addressing
once, so two things become possible:

  1. the watch agent can name the exact cells a source would change, instead of
     emitting a vague "suggested action";
  2. RegWatch can read the workbook back and know which field it is looking at.

Sortie : data/cellmap/<reg>.json - une carte par reglementation (voir tools/cellmap.py).

    python3 tools/excel_cellmap.py "ressources/CYBER WATCH5_Technical inventory.xlsx"
    python3 tools/excel_cellmap.py "<classeur.xlsx>" --reg rec        # un autre classeur
    python3 tools/excel_cellmap.py "<classeur.xlsx>" --check          # rien n'est ecrit

--check compare la carte enregistree au classeur et nomme ce qui a bouge :
feuille disparue, ligne d'en-tete deplacee, colonne renommee, pays deplace de
ligne. C'est le controle a lancer apres chaque mise a jour du classeur : sans
lui, un routage devenu faux ne se voit pas - la veille designerait la mauvaise
cellule sans qu'aucune erreur n'apparaisse. Code de sortie 1 en cas d'ecart.

Read-only: the workbook is never written back. Columns holding formulas are flagged
`writable: false` — those are computed by Excel and must never be overwritten.
"""

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from cellmap import DEFAULT_REG, path_for  # noqa: E402

try:
    import openpyxl
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent

# The sheets that hold one row per country. Hidden helper sheets are excluded.
SHEETS = [
    "ID - P1",
    "Incident reporting - P1",
    "Registration - P1",
    "Cybersecurity frameworks - P1",
    "Audit & Controls - P1",
    "Sanctions - P2",
    "Authority - P3",
]

# Columns that are working notes, not comparative data.
SKIP_HEADERS = {"repartition", "répartition", "column1"}

# Workbook country label -> RegWatch ISO code.
ISO = {
    "austria": "AT", "belgium": "BE", "bulgaria": "BG", "croatia": "HR",
    "cyprus": "CY", "czechia": "CZ", "czech republic": "CZ", "denmark": "DK",
    "estonia": "EE", "finland": "FI", "france": "FR", "germany": "DE",
    "greece": "GR", "hungary": "HU", "ireland": "IE", "italy": "IT",
    "latvia": "LV", "lithuania": "LT", "luxembourg": "LU", "malta": "MT",
    "netherlands": "NL", "norway": "NO", "poland": "PL", "portugal": "PT",
    "romania": "RO", "slovakia": "SK", "slovenia": "SI", "spain": "ES",
    "sweden": "SE", "united kingdom": "GB", "uk": "GB",
}

# Aggregate rows at the bottom of each sheet — not countries, expected, not a warning.
AGGREGATES = {"total", "total/moyenne", "moyenne", "average"}


def clean(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def find_header_row(ws):
    """The header row is the one carrying the 'Country' label (row 3 or 4 here)."""
    for r in range(1, 12):
        for cell in ws[r]:
            if clean(cell.value).lower() == "country":
                return r, cell.column_letter
    return None, None


def build(path):
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)

    sheets = {}
    for name in SHEETS:
        if name not in values.sheetnames:
            print("  skipped (absent): %s" % name)
            continue
        ws, wf = values[name], formulas[name]
        hrow, ccol = find_header_row(ws)
        if not hrow:
            print("  skipped (no Country column): %s" % name)
            continue

        # country -> row
        rows = {}
        for r in range(hrow + 1, ws.max_row + 1):
            label = clean(ws["%s%d" % (ccol, r)].value)
            iso = ISO.get(label.lower())
            if iso:
                rows[iso] = r
            elif label and label.lower() not in AGGREGATES:
                print("  %s: unmapped country label %r (row %d)" % (name, label, r))

        # field -> column, with a writability flag taken from the first data row
        fields = []
        probe = min(rows.values()) if rows else hrow + 1
        for cell in ws[hrow]:
            header = clean(cell.value)
            if not header or cell.column_letter == ccol:
                continue
            if header.lower() in SKIP_HEADERS or re.fullmatch(r"\d+", header):
                continue
            raw = wf["%s%d" % (cell.column_letter, probe)].value
            fields.append({
                "column": cell.column_letter,
                "label": header,
                "writable": not (isinstance(raw, str) and raw.startswith("=")),
            })

        sheets[name] = {"headerRow": hrow, "countryColumn": ccol, "rows": rows, "fields": fields}
        print("  %-32s %2d countries x %2d fields" % (name, len(rows), len(fields)))

    return sheets


def compare(stored, fresh):
    """Ce qui a bouge entre la carte enregistree et le classeur d'aujourd'hui."""
    drift = []
    for name, was in stored.items():
        now = fresh.get(name)
        if not now:
            drift.append("feuille absente du classeur : %s" % name)
            continue
        if was.get("headerRow") != now["headerRow"]:
            drift.append("%s : ligne d'en-tête %s -> %s" % (name, was.get("headerRow"), now["headerRow"]))
        if was.get("countryColumn") != now["countryColumn"]:
            drift.append("%s : colonne des pays %s -> %s" % (name, was.get("countryColumn"), now["countryColumn"]))
        by_now = {f["column"]: f["label"] for f in now["fields"]}
        for f in was.get("fields", []):
            label = by_now.get(f["column"])
            if label is None:
                drift.append("%s!%s : colonne disparue (était « %s »)" % (name, f["column"], f["label"]))
            elif label != f["label"]:
                drift.append("%s!%s : « %s » -> « %s »" % (name, f["column"], f["label"], label))
        for col, label in by_now.items():
            if col not in {f["column"] for f in was.get("fields", [])}:
                drift.append("%s!%s : nouvelle colonne « %s »" % (name, col, label))
        for iso, row in (was.get("rows") or {}).items():
            if now["rows"].get(iso) != row:
                drift.append("%s : %s ligne %s -> %s" % (name, iso, row, now["rows"].get(iso, "absent")))
    for name in fresh:
        if name not in stored:
            drift.append("nouvelle feuille : %s" % name)
    return drift


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--reg", default=DEFAULT_REG, help="réglementation (nis2 par défaut)")
    ap.add_argument("--check", action="store_true",
                    help="comparer la carte enregistrée au classeur, sans rien écrire")
    args = ap.parse_args()
    path = Path(args.workbook).expanduser()
    if not path.exists():
        raise SystemExit("classeur introuvable : %s" % path)

    sheets = build(path)

    out = path_for(args.reg)
    if args.check:
        if not out.exists():
            raise SystemExit("aucune carte enregistrée pour « %s » : lancer sans --check" % args.reg)
        stored = json.loads(out.read_text(encoding="utf-8"))
        drift = compare(stored.get("sheets", {}), sheets)
        if not drift:
            print("\ncarte à jour : %s correspond au classeur %s" % (out.relative_to(ROOT), path.name))
            return 0
        print("\n%d écart(s) entre %s et le classeur %s :" % (len(drift), out.relative_to(ROOT), path.name))
        for d in drift:
            print("  !! %s" % d)
        print("\nRelancer sans --check pour reconstruire la carte, puis vérifier le routage")
        print("de la file de veille (tools/veille_to_watchitems.py).")
        return 1

    total = sum(len(s["rows"]) * len(s["fields"]) for s in sheets.values())
    writable = sum(len(s["rows"]) * sum(1 for f in s["fields"] if f["writable"])
                   for s in sheets.values())

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({
        "workbook": path.name,
        "sheets": sheets,
        "addressableCells": total,
        "writableCells": writable,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print("\n%d addressable cells (%d writable) -> %s"
          % (total, writable, out.relative_to(ROOT)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
