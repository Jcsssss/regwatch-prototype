#!/usr/bin/env python3
"""Porter la feuille « Sectors - P3 » : le perimetre retenu par chaque pays.

    python3 tools/sectors_to_js.py <classeur.xlsx>   # -> src/reg/nis2/data_sectors.js

Pourquoi un extracteur a part
  La feuille n'a pas la forme des autres : deux lignes d'en-tete (le secteur sur
  la ligne 2, fusionnee, le sous-secteur sur la ligne 3), et pour chaque secteur
  un meme trio de colonnes - la couverture de chaque sous-secteur de la
  directive, les sous-secteurs ajoutes, les autres precisions. Lue ligne par
  ligne comme les autres feuilles, elle donnerait dix-sept colonnes toutes
  intitulees « Additional sub-sectors ».

Lecture des cases
  YES          le sous-secteur est couvert comme dans la directive
  NO           le pays s'en ecarte (la legende de la ligne 1 le dit)
  NC           non communique
  Ajouts et precisions : « NO » veut dire rien a signaler, et « YES : texte »
  est ramene au texte, seul porteur d'information.

La colonne « Repartition » (prenoms de l'equipe) est ecartee, comme dans
sheets_to_countries.py : la page est servie publiquement.
"""

import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
CELLMAP = ROOT / "data" / "excel-cellmap.json"
OUT = ROOT / "src" / "reg" / "nis2" / "data_sectors.js"
SHEET = "Sectors - P3"

# Secteur du classeur -> cle d'icone de la fiche (SECT_ICONS) et annexe.
SECTORS = [
    ("Energy sector", "energy", "ee"),
    ("Transportation sector", "transport", "ee"),
    ("Banking sector", "bank", "ee"),
    ("Financial markets infrastructures", "fmi", "ee"),
    ("Health", "health", "ee"),
    ("Wastewater", "waste", "ee"),
    ("Digital infrastructures", "digital", "ee"),
    ("ICT service management", "ict", "ee"),
    ("Public administration", "public", "ee"),
    ("Space", "space", "ee"),
    ("Postal and shipping services", "post", "ei"),
    ("Waste management", "garbage", "ei"),
    ("Manufacturing, production and distribution of chemical products", "chem", "ei"),
    ("Food production, processing, and distribution", "food", "ei"),
    ("Manufacturing", "manu", "ei"),
    ("Digital providers", "provider", "ei"),
    ("Research", "research", "ei"),
]
SKIP = re.compile(r"^(r[ée]partition|column\d*|colonne\d*)$", re.I)


def clean(v):
    if v is None:
        return None
    s = str(v).replace("_x000D_", " ").replace("\xa0", " ")
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n{3,}", "\n\n", s).strip()
    if not s or s.lower() in ("n/a", "na", "-", "--", "?", "."):
        return None
    if re.fullmatch(r"n\.?\s?c\.?|tbc|tbd", s, re.I):
        return "NC"
    return s


def flag(v):
    """Case de couverture : YES, NO ou NC."""
    v = clean(v)
    if v is None:
        return None
    if re.match(r"^yes\b", v, re.I):
        return "YES"
    if re.match(r"^no\b", v, re.I):
        return "NO"
    return "NC" if v == "NC" else v


def note(v):
    """Ajout ou precision : le texte, ou None s'il n'y a rien a signaler."""
    v = clean(v)
    if v is None or re.fullmatch(r"no\.?|non|none", v, re.I):
        return None
    if v == "NC":
        return "NC"
    v = re.sub(r"^yes\s*[:\-]\s*", "", v, flags=re.I).strip()
    return None if re.fullmatch(r"yes\.?", v, re.I) else v


def main():
    if len(sys.argv) < 2:
        raise SystemExit("usage: sectors_to_js.py <classeur.xlsx>")
    src = Path(sys.argv[1]).expanduser()
    cm = json.loads(CELLMAP.read_text(encoding="utf-8"))
    wb = load_workbook(src, data_only=True)
    ws = wb[SHEET]

    # Les pays sont reperes par leur nom, que la feuille ID - P1 associe a un
    # code dans la cartographie : la feuille des secteurs n'est pas cartographiee.
    idspec = cm["sheets"]["ID - P1"]
    ids = wb["ID - P1"]
    name2iso = {}
    for iso, row in idspec["rows"].items():
        n = clean(ids["%s%d" % (idspec["countryColumn"], row)].value)
        if n:
            name2iso[n.lower()] = iso

    # Colonnes de chaque secteur, lues dans la feuille : ligne 2 pour le
    # secteur (cellule fusionnee, donc renseignee sur sa premiere colonne
    # seulement), ligne 3 pour le sous-secteur.
    cols, cur = {}, None
    newcol = None
    for c in range(1, ws.max_column + 1):
        s = clean(ws.cell(2, c).value)
        if s:
            cur = s
        h = clean(ws.cell(3, c).value) or ""
        if not cur or SKIP.match(h):
            continue
        if cur.upper() == "NEW SECTORS":
            if re.match(r"addition", h, re.I):
                newcol = c
            continue
        bag = cols.setdefault(cur, {"subs": [], "add": None, "other": None})
        if re.match(r"additional", h, re.I):
            bag["add"] = c
        elif re.match(r"other ch", h, re.I):
            bag["other"] = c
        else:
            bag["subs"].append((c, h))

    known = {s[0] for s in SECTORS}
    missing = [n for n in known if n not in cols]
    extra = [n for n in cols if n not in known]
    if missing or extra:
        raise SystemExit("secteurs inattendus - manquants %s, nouveaux %s" % (missing, extra))

    sectors = []
    for name, key, annex in SECTORS:
        subs = [h for _, h in cols[name]["subs"]]
        sectors.append({"k": key, "name": name, "annex": annex, "subs": subs})

    rows, unknown = {}, []
    for r in range(4, ws.max_row + 1):
        n = clean(ws.cell(r, 2).value)
        if not n:
            continue
        iso = name2iso.get(n.lower())
        if not iso:
            unknown.append(n)
            continue
        out = {}
        for name, key, _ in SECTORS:
            spec = cols[name]
            cov = [flag(ws.cell(r, c).value) for c, _ in spec["subs"]]
            item = {"cov": cov}
            if spec["add"]:
                v = note(ws.cell(r, spec["add"]).value)
                if v: item["add"] = v
            if spec["other"]:
                v = note(ws.cell(r, spec["other"]).value)
                if v: item["other"] = v
            out[key] = item
        if newcol:
            v = note(ws.cell(r, newcol).value)
            if v: out["_new"] = v
        rows[iso] = out

    upd = clean(ws.cell(1, 1).value) or ""
    m = re.search(r"(\d{2})/(\d{2})/(\d{4})", upd)
    payload = {"updated": "%s-%s-%s" % (m.group(3), m.group(2), m.group(1)) if m else None,
               "sectors": sectors, "rows": rows}
    OUT.write_text(
        "/* ---- Perimetre sectoriel par pays (feuille « Sectors - P3 »).\n"
        "   Genere par tools/sectors_to_js.py - ne pas editer a la main. ---- */\n"
        "const SECTOR_DATA = %s;\n" % json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8")
    print("%d secteurs, %d pays (ignores : %s)" % (len(sectors), len(rows), ", ".join(unknown) or "aucun"))
    print("%d Ko -> %s" % (OUT.stat().st_size // 1024, OUT.relative_to(ROOT)))


if __name__ == "__main__":
    main()
