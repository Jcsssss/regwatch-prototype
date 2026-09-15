#!/usr/bin/env python3
"""Embarquer le detail des mesures de cybersecurite d'un pays.

    python3 tools/requirements_to_js.py
    -> src/reg/nis2/data_requirements.js

Lit ressources/NIS2_<Pays>_Requirements.xlsx, un fichier par pays. La fiche
pays affiche les compteurs du Cyber Watch - 25 themes, 293 exigences - et une
loupe ouvre ce qu'il y a derriere.

Ce que ces fichiers portent, et ce qu'ils ne portent pas
  Les mesures du referentiel national, une par ligne : type, reference legale,
  perimetre officiel, applicabilite aux entites essentielles et importantes.
  Ils donnent aussi les compteurs d'exigences, mais pas la liste des exigences
  elles-memes. La fiche n'invente donc aucune repartition des 293 exigences par
  mesure : elle montre les mesures, et explique les compteurs avec les textes
  du fichier.
"""

import json
import re
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "ressources"
OUT = ROOT / "src" / "reg" / "nis2" / "data_requirements.js"

ISO = {
    "Austria": "AT", "Belgium": "BE", "Bulgaria": "BG", "Croatia": "HR", "Cyprus": "CY",
    "Czech Republic": "CZ", "Czechia": "CZ", "Denmark": "DK", "Estonia": "EE",
    "Finland": "FI", "France": "FR", "Germany": "DE", "Greece": "GR", "Hungary": "HU",
    "Ireland": "IE", "Italy": "IT", "Latvia": "LV", "Lithuania": "LT",
    "Luxembourg": "LU", "Malta": "MT", "Netherlands": "NL", "Norway": "NO",
    "Poland": "PL", "Portugal": "PT", "Romania": "RO", "Slovakia": "SK",
    "Slovenia": "SI", "Spain": "ES", "Sweden": "SE", "United Kingdom": "GB",
}


def txt(v):
    return re.sub(r"[ \t]+", " ", str(v).replace("\xa0", " ")).strip() if v is not None else ""


def cells(row):
    return [txt(c) for c in row if c not in (None, "")]


def parse(path):
    name = re.match(r"NIS2_(.+)_Requirements\.xlsx$", path.name).group(1)
    iso = ISO.get(name)
    if not iso:
        return None, None
    wb = load_workbook(path, read_only=True, data_only=True)

    # --- les mesures ---
    rows = [cells(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    rows = [r for r in rows if r]
    note = next((r[0] for r in rows if len(r) == 1 and len(r[0]) > 120), "")
    hi = next(i for i, r in enumerate(rows) if "ID" in r and "Reference" in r)
    raw_hdr = [txt(c) for c in list(wb.worksheets[0].iter_rows(values_only=True))[0]]
    # l'en-tete est lu ligne pour ligne, colonnes vides comprises
    grid = [[txt(c) for c in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
    hrow = next(r for r in grid if "ID" in r and "Reference" in r)
    col = {h: i for i, h in enumerate(hrow) if h}
    measures = []
    for r in grid:
        rid = r[col["ID"]] if col.get("ID") is not None and len(r) > col["ID"] else ""
        if not re.match(r"^[A-Z]{2}-\d+", rid):
            continue
        get = lambda k: r[col[k]] if k in col and len(r) > col[k] else ""
        ie_note = get("Applicable to IE")
        measures.append({
            "id": rid, "type": get("Measure type"), "name": get("Measure (theme)"),
            "ref": get("Reference"), "scope": get("Requirement / official scope"),
            "ee": get("EE (higher regime)").upper() == "X",
            "ie": get("IE (lower regime)").upper() == "X",
            "ieNote": "" if ie_note.lower() in ("yes", "") else ie_note,
            "itot": get("IT / OT applicability"),
        })

    # --- les compteurs et les textes qui les expliquent ---
    kpi = [cells(r) for r in wb.worksheets[1].iter_rows(values_only=True)] if len(wb.worksheets) > 1 else []
    kpi = [r for r in kpi if r]
    counts = {}
    for i, r in enumerate(kpi):
        if r and r[0].startswith("Number of cyber themes for EE") and i + 1 < len(kpi):
            vals = kpi[i + 1]
            keys = ["themesEE", "themesIE", "reqEE", "reqIE"]
            counts = {k: int(float(v)) for k, v in zip(keys, vals) if re.match(r"^\d+(\.0)?$", v)}
            break
    flat = [" ".join(r) for r in kpi]
    def after(label):
        for i, line in enumerate(flat):
            if line.strip().upper() == label and i + 1 < len(flat):
                return flat[i + 1]
        return ""
    texts = {
        "intro": next((l for l in flat if l.startswith("The ") and "differentiation" in l), ""),
        "themes": next((l for l in flat if l.startswith("Themes:")), ""),
        "req": next((l for l in flat if l.startswith("Requirements:")), ""),
        "caveat": after("IMPORTANT CAVEAT"),
    }
    src_line = after("SOURCES")
    sources = [{"label": txt(m.group(1)).strip(" |:"), "url": m.group(2)}
               for m in re.finditer(r"([^|]*?):\s*(https?://\S+)", src_line)]
    return iso, {"country": name, "note": note, "counts": counts, "texts": texts,
                 "sources": sources, "measures": measures}


def main():
    out = {}
    for f in sorted(SRC.glob("NIS2_*_Requirements.xlsx")):
        iso, data = parse(f)
        if not iso:
            print("   %s : pays non reconnu, ecarte" % f.name)
            continue
        out[iso] = data
        ie = sum(1 for m in data["measures"] if m["ie"])
        print("   %s  %d mesures (%d applicables EI) | compteurs %s | %d sources"
              % (iso, len(data["measures"]), ie, data["counts"], len(data["sources"])))
    OUT.write_text(
        "/* ---- Detail des mesures de cybersecurite par pays, derriere les compteurs\n"
        "   de la fiche. Genere par tools/requirements_to_js.py - ne pas editer. ---- */\n"
        "const FW_REQUIREMENTS = %s;\n" % json.dumps(out, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8")
    print("%d pays -> %s (%d Ko)" % (len(out), OUT.relative_to(ROOT), OUT.stat().st_size // 1024))


if __name__ == "__main__":
    main()
