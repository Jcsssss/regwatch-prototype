#!/usr/bin/env python3
"""Embarquer dans l'outil les correspondances validees entre referentiels.

    python3 tools/mapping_to_js.py
    -> src/reg/nis2/data_mapping.js

D'ou viennent les donnees
  De l'agent de correspondance (ressources/Agent-Mapping-NIS2-main), qui
  compare deux referentiels exigence par exigence et produit un classeur. Le
  moteur ne tourne pas dans RegWatch : il lui faut un backend, Azure OpenAI et
  des minutes de calcul par paire, la ou l'outil est un fichier qui s'ouvre hors
  ligne. On embarque ses resultats, comme on embarque le Cyber Watch.

Seulement les classeurs valides
  L'agent produit un classeur a chaque execution, parfois six pour la meme paire.
  Seuls ceux dont le nom commence par « Validated_ » ont ete relus : ce sont les
  seuls embarques. C'est la regle de la file de veille et des phrases generees -
  rien de non relu ne rejoint l'outil.

Le poids
  Une paire pese environ 660 Ko au niveau des exigences, parce que les textes
  sont longs et repetes d'une ligne a l'autre. Deux mesures : chaque texte
  d'exigence n'est ecrit qu'une fois, reference par son identifiant ; et le
  detail est compresse, decompresse par la page au premier affichage seulement.
  Le resume - taux, repartition, categories - reste en clair : la fiche pays le
  lit sans rien decompresser.
"""

import base64
import gzip
import json
import re
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    raise SystemExit("pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "ressources" / "Agent-Mapping-NIS2-main" / "output"
OUT = ROOT / "src" / "reg" / "nis2" / "data_mapping.js"

# Ce que le nom du fichier ne dit pas : le nom d'affichage de chaque referentiel
# et le pays qu'il concerne, pour relier la correspondance a sa fiche.
#
# « France 2.3 » dans les feuilles du classeur valide designe en realite ReCyF
# 2.5 : le cote francais compte exactement 152 exigences, le nombre de mesures
# de ReCyF 2.5, et le fichier a ete nomme ainsi a la validation. Le libelle 2.3
# est celui de la bibliotheque de l'agent au moment de l'execution.
FRAMEWORKS = {
    "Cyfun 2025": {"name": "CyFun 2025", "iso": "BE"},
    "France 2.3": {"name": "ReCyF 2.5", "iso": "FR"},
    "France 2.5": {"name": "ReCyF 2.5", "iso": "FR"},
}

REL = {  # libelle du classeur -> code court, dans l'ordre de couverture
    "Fully covered": "full", "Largely covered": "large", "Partially covered": "partial",
    "Indirectly covered": "indirect", "Not covered": "none",
}


def norm(s):
    return re.sub(r"[ \t]+", " ", str(s).replace("\xa0", " ")).strip() if s is not None else ""


def table(ws, first_col_label):
    """Les lignes d'un tableau, a partir de la ligne d'en-tete qui porte ce libelle."""
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows)
              if r and first_col_label in [norm(x) for x in r])
    hdr = [norm(x) for x in rows[hi]]
    return [dict(zip(hdr, r)) for r in rows[hi + 1:] if r and any(v not in (None, "") for v in r)]


def fw(label):
    info = FRAMEWORKS.get(norm(label))
    return info or {"name": norm(label), "iso": None}


def slug(s):
    return re.sub(r"[^a-z0-9]+", "", s.lower())


def main():
    files = sorted(SRC.glob("Validated_*.xlsx"))
    if not files:
        raise SystemExit("aucun classeur valide dans %s" % SRC)

    index, detail = {}, {}
    for f in files:
        wb = load_workbook(f, read_only=True, data_only=True)
        directions = [n for n in wb.sheetnames if "->" in n and not n.startswith("Atomic")]
        if not directions:
            print("   %s : aucune feuille de correspondance, ecarte" % f.name)
            continue

        texts = {}        # referentiel -> { id: texte }
        dirs = []
        pair_names = None
        for sheet in directions:
            rows = [r for r in table(wb[sheet], "Source control ID") if r.get("Source control ID")]
            if not rows:
                continue
            a, b = fw(rows[0]["Source regulation"]), fw(rows[0]["Target regulation"])
            pair_names = pair_names or sorted([a["name"], b["name"]])
            ta = texts.setdefault(a["name"], {})
            out, counts, total = [], {k: 0 for k in REL.values()}, 0
            for r in rows:
                sid = norm(r["Source control ID"])
                ta[sid] = norm(r.get("Source requirement"))
                rel = REL.get(norm(r.get("Coverage relationship")), "none")
                counts[rel] += 1
                lvl = r.get("Coverage level")
                lvl = int(lvl) if isinstance(lvl, (int, float)) else 0
                total += lvl
                tids = [t for t in re.split(r"[,;\n]+", norm(r.get("Target control ID(s)"))) if t.strip()]
                out.append({
                    "id": sid, "cat": norm(r.get("ENISA category")), "to": [t.strip() for t in tids],
                    "rel": rel, "lvl": lvl, "gap": norm(r.get("Gap")),
                    "det": norm(r.get("Detailed gap")), "act": norm(r.get("Action plan")),
                    "pri": norm(r.get("Review priority")),
                })
            dirs.append({"from": a, "to": b, "n": len(out), "avg": round(total / max(1, len(out)), 1),
                         "counts": counts, "rows": out})

        pid = "-".join(slug(n) for n in pair_names)
        # la repartition par categorie ENISA, par sens, calculee sur les lignes
        # plutot que lue dans l'onglet de synthese, qui agrege les deux sens
        summary = []
        for d in dirs:
            cats = {}
            for r in d["rows"]:
                c = cats.setdefault(r["cat"] or "-", {"n": 0, "sum": 0, "gaps": 0})
                c["n"] += 1
                c["sum"] += r["lvl"]
                c["gaps"] += 1 if r["rel"] == "none" else 0
            summary.append({
                "from": d["from"], "to": d["to"], "n": d["n"], "avg": d["avg"], "counts": d["counts"],
                "cats": sorted(({"cat": k, "n": v["n"], "avg": round(v["sum"] / v["n"], 1), "gaps": v["gaps"]}
                                for k, v in cats.items()), key=lambda x: -x["n"]),
            })
        index[pid] = {"file": f.name, "names": pair_names, "validated": True, "dirs": summary}
        detail[pid] = {"texts": texts, "rows": [d["rows"] for d in dirs]}
        print("   %-28s %s" % (pid, " | ".join("%s -> %s : %d exigences, %.1f %% en moyenne"
              % (d["from"]["name"], d["to"]["name"], d["n"], d["avg"]) for d in summary)))

    raw = json.dumps(detail, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    gz = base64.b64encode(gzip.compress(raw, 9)).decode("ascii")
    OUT.write_text(
        "/* ---- Correspondances validees entre referentiels, issues de l'agent de\n"
        "   correspondance. Genere par tools/mapping_to_js.py - ne pas editer.\n"
        "   MAPPING_INDEX est en clair ; le detail est compresse (MAPPING_GZ). ---- */\n"
        "const MAPPING_INDEX = %s;\n"
        "const MAPPING_GZ = \"%s\";\n" % (json.dumps(index, ensure_ascii=False, separators=(",", ":")), gz),
        encoding="utf-8")
    print("\n%d correspondance(s) ; detail %d Ko brut -> %d Ko compresse ; fichier %d Ko"
          % (len(index), len(raw) // 1024, len(gz) // 1024, OUT.stat().st_size // 1024))


if __name__ == "__main__":
    main()
