#!/usr/bin/env python3
"""Extraire les themes du referentiel de chaque pays depuis CYBER WATCH 2.

    python3 tools/fw_themes_from_pptx.py "ressources/CYBER WATCH2_....pptx"
    -> src/reg/nis2/data_fwthemes.js

Pourquoi lire un PowerPoint
  Les themes du referentiel national - leur nom, leur numero, le nombre
  d'exigences et la phrase qui les decrit - n'existent nulle part ailleurs. Le
  classeur ne porte que des compteurs : « 17 themes, 218 mesures ». Les intitules
  eux-memes ont ete rediges dans le support, et c'est la seule source.

  Le generateur de rapport les reprend donc tels quels. L'alternative etait de
  les fabriquer a partir de la grille de couverture, ce qui produisait une slide
  annoncant « 6 themes » puis en affichant quinze : deux comptages differents
  pour la meme question, sur la meme slide.

Comment les blocs sont retrouves
  Une slide n'a pas de structure, seulement des rectangles places. On les trie
  par colonne (abscisse) puis par hauteur, et on distingue trois formes :
  l'intitule d'une famille (texte court, sans numero), un theme (« 4 - Nom (5) »
  suivi de sa description), et le reste, ignore.
"""

import json
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:
    raise SystemExit("pip install python-pptx")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "src" / "reg" / "nis2" / "data_fwthemes.js"

# « 12 - Security incident management (7) » ou « 1 – IS inventory (3) »
THEME = re.compile(r"^\s*(\d{1,2})\s*[-–—]\s*(.+?)\s*(?:\((\d+)\)\s*)?$")
TITLE = re.compile(r"Summary of themes of NIS\s*2\s*[-–—]\s*(.+?)\s*$", re.I)
INTRO = re.compile(r"(\d+)\s+themes?.{0,40}?(\d+)\s+security measures.*?"
                   r"(?:(\d+)\s+security measures)?", re.I | re.S)

NAMES = {
 "AUSTRIA":"AT","BELGIUM":"BE","BULGARIA":"BG","CROATIA":"HR","CYPRUS":"CY",
 "CZECH REPUBLIC":"CZ","CZECHIA":"CZ","DENMARK":"DK","ESTONIA":"EE","FINLAND":"FI",
 "FRANCE":"FR","GERMANY":"DE","GREECE":"GR","HUNGARY":"HU","IRELAND":"IE",
 "ITALY":"IT","ITA":"IT","LATVIA":"LV","LITHUANIA":"LT","LUXEMBOURG":"LU",
 "MALTA":"MT","NETHERLANDS":"NL","NORWAY":"NO","POLAND":"PL","PORTUGAL":"PT",
 "ROMANIA":"RO","SLOVAKIA":"SK","SLOVENIA":"SI","SPAIN":"ES","SWEDEN":"SE",
 "UNITED KINGDOM":"GB",
}


def blocks(shapes, acc):
    for sh in shapes:
        try:
            if sh.shape_type == 6 and hasattr(sh, "shapes"):
                blocks(sh.shapes, acc)
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                acc.append((int(sh.left or 0), int(sh.top or 0),
                            sh.text_frame.text.strip()))
        except Exception:
            pass
    return acc


def clean(s):
    s = s.replace("’", "'").replace("\xa0", " ")
    return re.sub(r"[ \t]+", " ", s).strip()


def parse(slide):
    """Un pays et ses familles, ou None si la slide n'est pas une slide themes."""
    items = blocks(slide.shapes, [])
    title = next((t for _, _, t in items if TITLE.search(t)), None)
    if not title:
        return None
    raw = TITLE.search(clean(title)).group(1)
    fw = ""
    m = re.search(r"\(([^)]+)\)\s*$", raw)
    if m:
        fw = m.group(1).strip()
        raw = raw[:m.start()].strip()
    iso = NAMES.get(raw.upper())
    if not iso:
        return None

    intro = next((clean(t) for _, _, t in items
                  if "security measures" in t.lower() and "themes" in t.lower()), "")
    nums = re.findall(r"(\d+)\s+themes?|(\d+)\s+security measures", intro, re.I)
    flat = [a or b for a, b in nums]
    themes_n = int(flat[0]) if len(flat) > 0 else None
    ee = int(flat[1]) if len(flat) > 1 else None
    ie = int(flat[2]) if len(flat) > 2 else None

    # colonnes : on regroupe les abscisses proches, puis on lit de haut en bas
    cand = [(x, y, clean(t)) for x, y, t in items
            if t.strip() not in ("IE", "(x)") and not TITLE.search(t)
            and "security measures" not in t.lower()]
    cand.sort(key=lambda v: (round(v[0] / 914400.0 * 2), v[1]))

    fams, cur = [], None
    for x, y, txt in cand:
        first, _, rest = txt.partition("\n")
        m = THEME.match(first)
        if m and rest.strip():
            if cur is None:
                cur = {"name": "", "themes": []}
                fams.append(cur)
            cur["themes"].append({
                "n": int(m.group(1)),
                "name": m.group(2).strip(),
                "req": int(m.group(3)) if m.group(3) else None,
                "desc": " ".join(l.strip() for l in rest.split("\n") if l.strip()),
            })
        elif len(txt) < 46 and "\n" not in txt and not m:
            # un intitule de famille : court, sur une ligne, sans numero
            cur = {"name": txt, "themes": []}
            fams.append(cur)
    fams = [f for f in fams if f["themes"]]
    if not fams:
        return None
    return iso, {"fw": fw, "themes": themes_n, "ee": ee, "ie": ie, "families": fams}


def main():
    if len(sys.argv) < 2:
        raise SystemExit("usage: fw_themes_from_pptx.py <CYBER WATCH2 ....pptx>")
    src = Path(sys.argv[1]).expanduser()
    if not src.exists():
        raise SystemExit("support introuvable : %s" % src)

    out, skipped = {}, 0
    for slide in Presentation(src).slides:
        got = parse(slide)
        if not got:
            skipped += 1
            continue
        iso, data = got
        # Plusieurs slides par pays quand il existe plusieurs referentiels. On
        # garde celle qui porte la phrase de synthese chiffree : c'est la slide
        # aboutie, les autres sont des variantes de travail. « La plus fournie »
        # choisissait a l'inverse pour la Belgique une version sans compteurs.
        prev = out.get(iso)
        def score(d):
            return (1 if d["themes"] else 0,
                    sum(1 for f in d["families"] for t in f["themes"] if t["req"]),
                    sum(len(f["themes"]) for f in d["families"]))
        if prev is None or score(data) > score(prev):
            out[iso] = data

    OUT.write_text(
        "/* ---- Themes du referentiel national, par pays.\n"
        "   Genere par tools/fw_themes_from_pptx.py depuis CYBER WATCH 2.\n"
        "   Ne pas editer a la main : le support est la source. ---- */\n"
        "const FW_THEMES = %s;\n" % json.dumps(out, ensure_ascii=False,
                                               separators=(",", ":")),
        encoding="utf-8")
    total = sum(sum(len(f["themes"]) for f in v["families"]) for v in out.values())
    print("%d pays, %d themes (%d slides ecartees)" % (len(out), total, skipped))
    for iso in sorted(out):
        v = out[iso]
        print("   %s  %-28s %d familles, %d themes" %
              (iso, (v["fw"] or "-")[:28], len(v["families"]),
               sum(len(f["themes"]) for f in v["families"])))
    print("%d Ko -> %s" % (OUT.stat().st_size // 1024, OUT.relative_to(ROOT)))


if __name__ == "__main__":
    main()
