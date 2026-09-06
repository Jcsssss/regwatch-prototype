#!/usr/bin/env python3
"""La présentation RegWatch en trois slides, au format PowerPoint.

    python3 tools/make_deck_pptx.py            # -> RegWatch - presentation.pptx

Mêmes chiffres, mêmes couleurs et même composition que la version web, mais
livrée dans le format où elle sera réellement montrée : un .pptx que le
consultant ouvre, retouche et intègre à son propre deck.

Deux choix qui comptent pour l'usage :

  le graphique est natif    PowerPoint le reconnaît comme un graphique, avec ses
                            données modifiables. Une image aurait été plus simple
                            à produire et inutilisable en réunion.
  la police est Aptos       la charte du cabinet, et la police par défaut de
                            Microsoft 365 : elle sera présente sur le poste pro.

Tous les chiffres sont repris du jeu de données de l'outil ; ils sont écrits ici
en clair pour que le deck reste lisible sans exécuter le pipeline, et la source
de chacun est indiquée en commentaire.
"""

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError:
    raise SystemExit("pip install python-pptx")

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "RegWatch - presentation.pptx"

# Charte, reprise de src/shell_top.html sans altération.
VIOLET = RGBColor(0x45, 0x1D, 0xC7)
VIOLET_DEEP = RGBColor(0x25, 0x0F, 0x6B)
VIOLET_SOFT = RGBColor(0xEC, 0xE7, 0xFB)
GREEN = RGBColor(0x04, 0xF0, 0x6A)
INK = RGBColor(0x16, 0x12, 0x2E)
INK2 = RGBColor(0x41, 0x3B, 0x60)
MUTED = RGBColor(0x6F, 0x6B, 0x8C)
LINE = RGBColor(0xE4, 0xE1, 0xF0)
SURF2 = RGBColor(0xF0, 0xEE, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN = RGBColor(0x7A, 0x54, 0x00)

UI = "Aptos"
W, H = 13.333, 7.5
RAIL = 1.0          # largeur du rail gauche, en pouces
X = 1.75            # marge de texte
TXTW = W - X - 0.9


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False,
         space=None, caps=False, align=PP_ALIGN.LEFT, spacing=1.3,
         anchor=MSO_ANCHOR.TOP):
    """`runs` : une chaîne, ou une liste de (texte, couleur, gras) par paragraphe.

    Une liste de listes fait plusieurs paragraphes ; une liste de tuples fait
    plusieurs runs dans un seul paragraphe, ce qui permet de colorer un membre
    de phrase sans casser le retour à la ligne.
    """
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    paras = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        chunks = para if isinstance(para, list) else [(para, color, bold)]
        for chunk in chunks:
            content, col, bd = chunk if isinstance(chunk, tuple) else (chunk, color, bold)
            r = p.add_run()
            r.text = content.upper() if caps else content
            r.font.size = Pt(size)
            r.font.name = UI
            r.font.bold = bd
            r.font.color.rgb = col
            if space:
                # L'interlettrage des intitulés : 100 = 1 point.
                r.font._rPr.set("spc", str(int(space * 100)))
    return tb


def rail(slide, number, label):
    box(slide, 0, 0, RAIL, H, fill=SURF2)
    box(slide, RAIL, 0, 0.008, H, fill=LINE)
    text(slide, 0, 0.5, RAIL, 0.4, number, size=15, bold=True, color=VIOLET,
         align=PP_ALIGN.CENTER)
    # Le libellé vertical : le rail porte le rang de la slide, qui est l'ordre
    # réel de lecture, pas un ornement.
    tb = text(slide, RAIL / 2 - 1.6, H / 2 - 0.25, 3.2, 0.5, label, size=8.5,
              color=MUTED, caps=True, space=2.4, align=PP_ALIGN.CENTER)
    tb.rotation = 270
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(RAIL / 2 - 0.035),
                                 Inches(H - 0.62), Inches(0.07), Inches(0.07))
    dot.fill.solid()
    dot.fill.fore_color.rgb = VIOLET
    dot.line.fill.background()
    dot.shadow.inherit = False


def header(slide, eyebrow, title_lines):
    text(slide, X, 0.62, TXTW, 0.25, eyebrow, size=9, color=MUTED, caps=True, space=2)
    text(slide, X, 0.98, TXTW, 1.5, title_lines, size=33, bold=True, spacing=1.06)


def slide_one(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, W, H, fill=WHITE)
    rail(s, "01", "Ce que c'est")
    header(s, "Le livrable", [
        [("Un fichier. Aucune installation.", INK, True)],
        [("Toute la transposition NIS 2.", VIOLET, True)],
    ])
    text(s, X, 2.62, 7.1, 1.6,
         "Le suivi de la transposition vivait dans un classeur Excel que personne "
         "ne pouvait ouvrir à deux. RegWatch en fait une page unique, consultable "
         "hors ligne, que le consultant envoie par mail ou ouvre chez le client "
         "sans rien installer — et qui reste la source d'où repart l'Excel comparatif.",
         size=14, color=INK2, spacing=1.45)

    box(s, X, 5.02, TXTW, 0.008, fill=LINE)
    figs = [
        # 29 fiches = 27 UE + GB + NO (src/reg/nis2/data_c1..c4.js)
        ("29", "fiches pays : les 27 États membres, plus le Royaume-Uni et la Norvège en comparateurs"),
        ("116", "textes officiels référencés, chacun avec son lien de publication"),  # data_docs.js
        ("210", "thèmes cyber cartographiés en 21 familles"),                          # data_themes.js
        ("2", "langues d'interface, français et anglais, sur l'intégralité du contenu"),
    ]
    colw = TXTW / 4
    for i, (value, label) in enumerate(figs):
        cx = X + i * colw
        if i:
            box(s, cx - 0.16, 5.28, 0.008, 1.25, fill=LINE)
        text(s, cx, 5.3, colw - 0.4, 0.7, value, size=34, bold=True, color=VIOLET_DEEP)
        text(s, cx, 6.03, colw - 0.4, 0.9, label, size=9.5, color=INK2, spacing=1.25)


def slide_two(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, W, H, fill=WHITE)
    rail(s, "02", "Ce qu'il montre")
    header(s, "Le fond réglementaire", [
        [("La date limite était le ", INK, True), ("17 octobre 2024", VIOLET, True), (".", INK, True)],
        [("Cinq États l'ont tenue.", INK, True)],
    ])

    # Retard de transposition des 23 États membres ayant légiféré, en mois.
    # Distribution mesurée sur data_c1..c4.js : 0->5, 1-6->5, 7-12->5, 13-18->7, 19+->1
    data = CategoryChartData()
    data.categories = ["0", "1–6", "7–12", "13–18", "19+"]
    data.add_series("États membres", (5, 5, 5, 7, 1))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(X - 0.15),
                            Inches(2.7), Inches(6.3), Inches(3.5), data)
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.vary_by_categories = False
    ser = plot.series[0]
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = VIOLET
    ser.format.line.fill.background()
    # Le respect du délai est un état, pas une valeur : il porte la seule
    # couleur sémantique de la charte.
    pt = ser.points[0]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = GREEN
    pt.format.line.fill.background()

    plot.has_data_labels = True
    dl = plot.data_labels
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(12)
    dl.font.bold = True
    dl.font.name = UI
    dl.font.color.rgb = INK

    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.maximum_scale = 8.0
    ca = ch.category_axis
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = RGBColor(0xCF, 0xCA, 0xE1)
    ca.tick_labels.font.size = Pt(11)
    ca.tick_labels.font.name = UI
    ca.tick_labels.font.color.rgb = MUTED

    text(s, X, 6.32, 6.0, 0.6,
         "Retard de transposition, en mois, pour les 23 États membres qui ont "
         "légiféré. En vert, ceux qui ont respecté le délai.",
         size=9.5, color=MUTED, spacing=1.3)

    stats = [
        (GREEN, "5 / 27", [("dans les délais : ", INK2, False),
                           ("Belgique, Croatie, Italie, Lettonie, Lituanie", INK, True)]),
        (WARN, "12 mois", [("de retard médian pour les 18 États ayant transposé "
                            "hors délai, jusqu'à 19 mois", INK2, False)]),
        (VIOLET, "4 / 27", [("encore sans loi : ", INK2, False),
                            ("France, Pays-Bas, Espagne, Irlande", INK, True),
                            (" — les quatre visés par le renvoi de la Commission "
                             "devant la Cour", INK2, False)]),
    ]
    sx, sy = 8.5, 2.75
    for color, value, label in stats:
        box(s, sx, sy, 0.035, 1.05, fill=color)
        text(s, sx + 0.22, sy - 0.04, 3.9, 0.45, value, size=22, bold=True, color=INK)
        text(s, sx + 0.22, sy + 0.42, 3.9, 0.75, [label], size=9.5, spacing=1.3)
        sy += 1.35


def slide_three(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, W, H, fill=WHITE)
    rail(s, "03", "Comment il tourne")
    header(s, "La chaîne, de la source au livrable", [
        [("Une veille automatisée, ", INK, True), ("validée à la main", VIOLET, True), (".", INK, True)],
    ])

    steps = [
        ("Étape 1 · Collecte", "L'agent lit les autorités",
         "Les autorités nationales de cybersécurité sont sondées chaque jour, "
         "article intégral récupéré quand la source l'autorise, dans sa langue "
         "d'origine.",
         "29 autorités · 22 flux vérifiés"),
        ("Étape 2 · Validation", "Un consultant tranche",
         "Chaque élément arrive dans une boîte de validation avec la phrase de la "
         "source qui a déclenché son classement, citée telle que publiée. Rien "
         "n'entre dans une fiche pays sans accord humain.",
         "89 éléments en file"),
        ("Étape 3 · Restitution", "Fiches, KPI, Excel",
         "Fiches pays comparables, indicateurs croisables à la demande, et un "
         "export Excel qui embarque les graphiques — pas seulement les données.",
         "Export .xlsx natif"),
    ]
    colw = TXTW / 3
    for i, (eyebrow, title, para, tag) in enumerate(steps):
        cx = X + i * colw
        if i:
            box(s, cx - 0.22, 2.35, 0.008, 2.35, fill=LINE)
        text(s, cx, 2.35, colw - 0.5, 0.25, eyebrow, size=9, color=VIOLET, caps=True, space=1.6)
        text(s, cx, 2.68, colw - 0.5, 0.4, title, size=17, bold=True, color=INK)
        text(s, cx, 3.18, colw - 0.5, 1.4, para, size=10.5, color=INK2, spacing=1.4)
        text(s, cx, 4.46, colw - 0.5, 0.3, tag, size=9, color=MUTED)

    box(s, X, 5.05, TXTW, 1.55, fill=VIOLET_SOFT)
    text(s, X + 0.35, 5.3, TXTW - 0.7, 0.8,
         [[("Un assistant répond aux questions du consultant", INK, True),
           (" — délais de transposition, autorités compétentes, organismes d'audit — "
            "uniquement à partir des données de l'outil, en citant ses sources. "
            "Il ne répond pas quand la donnée n'y est pas.", INK2, False)]],
         size=11, spacing=1.4)

    px = X + 0.35
    for label, live in [("NIS 2 · 29 pays", True), ("REC · 27 pays", True),
                        ("DORA · à venir", False)]:
        w = 0.13 * len(label) + 0.3
        # Une pastille, pas un champ de saisie : le rectangle droit se lisait
        # comme un formulaire sur le fond violet clair.
        pill = box(s, px, 6.06, w, 0.3, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                   fill=WHITE if live else None,
                   line=None if live else RGBColor(0xCF, 0xCA, 0xE1))
        pill.adjustments[0] = 0.5
        text(s, px, 6.13, w, 0.2, label, size=9,
             color=VIOLET_DEEP if live else MUTED, bold=live, align=PP_ALIGN.CENTER)
        px += w + 0.15


def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide_one(prs)
    slide_two(prs)
    slide_three(prs)
    prs.save(OUT)
    print("%d slides -> %s" % (len(prs.slides.__iter__.__self__._sldIdLst), OUT.name))
    print("%.0f Ko" % (OUT.stat().st_size / 1024))
    return 0


if __name__ == "__main__":
    sys.exit(main())
